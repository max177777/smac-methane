"""
Extracts Maryland's OFFICIAL state GHG inventory (MDE) methane figures from the
published workbooks, and folds them into the tool two ways:

  1. data/MD_official_inventory_ch4.csv — tidy CH4-by-category-by-year table.
  2. appended chunks in data/rag_corpus.json — so the Chat can cite MDE's own
     numbers when someone asks about Maryland.

WHY THIS MATTERS
Everything else in this tool is Climate TRACE — modeled from satellite and
activity data. This is Maryland's own bottom-up inventory, compiled by the
state. The two disagree, and that disagreement is itself the useful finding:
the Maryland Methane Action Plan already in our library notes MDE showing
statewide methane roughly flat while Climate TRACE shows it rising. Having
both in the chat lets that comparison be made with real numbers on each side
instead of hand-waving.

    python -m pip install pandas openpyxl
    python scripts/build_md_inventory.py

The Summary sheet in each workbook has a consistent shape: an indented tree of
category -> sub-category -> gas rows, with 20-year and 100-year GWP columns.
We walk it, tracking the current top-level category and sub-category, and keep
only the CH4 rows.
"""

from pathlib import Path
import json
import re
import sys

import pandas as pd

try:
    import openpyxl  # noqa: F401
except ImportError:
    sys.exit("Missing dependency. Run: python -m pip install pandas openpyxl")

UPLOAD_DIR = Path("md_inventory_source")   # put the five MD_*.xlsx files here
OUT_CSV = Path(__file__).parent.parent / "data" / "MD_official_inventory_ch4.csv"
RAG_PATH = Path(__file__).parent.parent / "data" / "rag_corpus.json"

# MMTCO2e -> tonnes CH4.  The workbook reports CO2-equivalents; to compare with
# Climate TRACE (which we display as tonnes of CH4) we divide back out by the
# GWP the workbook itself used, then convert million-tonnes to tonnes.
GWP100_IN_WORKBOOK = 28.0     # AR5 100-yr, per the workbook's own GWP table
MMT_TO_TONNES = 1_000_000

# Top-level categories in the Summary tree. Rows whose first non-empty cell is
# one of these starts a new category block.
TOP_LEVEL = {
    "Energy Use (CO2, CH4, N2O)": "Energy Use",
    "Electricity Use (Consumption Basis)": "Electricity Use",
    "Residential/Commercial/Industrial Fuel Use": "Buildings & Industry Fuel Use",
    "Transportation": "Transportation",
    "Fossil Fuel Industry": "Fossil Fuel Industry",
    "Industrial Processes and Product Use": "Industrial Processes",
    "Agriculture": "Agriculture",
    "Waste Management": "Waste Management",
    "Forestry and Land Use": "Forestry & Land Use",
}


def _clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def parse_workbook(path: Path, year: int) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = "Summary" if "Summary" in wb.sheetnames else wb.sheetnames[0]
    rows = list(wb[sheet].iter_rows(values_only=True))

    records = []
    category = None
    subcategory = None

    # Layout (verified against all five workbooks): column A is blank,
    # column B holds top-level categories, column C holds BOTH sub-categories
    # and gas rows, and columns D/E are the 20-yr and 100-yr GWP values.
    GAS_LABELS = {"CO2", "CH4", "N2O", "HFC, PFC, SF6", "HFCS & PFCS", "SF6"}

    for r in rows:
        row = list(r) + [None] * 6
        col_b, col_c = _clean(row[1]), _clean(row[2])
        gwp20 = row[3] if isinstance(row[3], (int, float)) else None
        gwp100 = row[4] if isinstance(row[4], (int, float)) else None

        # a new top-level category (column B). Only KNOWN categories count:
        # the sheet's title cell ("Maryland 2020 GHG Emissions by Sector") also
        # sits in column B, and accepting it as a category made the parser read
        # the GWP reference table right below it (CH4 | 84.0 | 28.0) as if it
        # were emissions — yielding a bogus 28/28*1e6 = 1,000,000 t entry.
        if col_b:
            matched_top = next((v for k, v in TOP_LEVEL.items()
                                if col_b.startswith(k[:18])), None)
            category = matched_top          # None if unrecognised -> rows skipped
            subcategory = None
            continue

        if not col_c:
            continue

        # column C is either a sub-category heading or a gas row
        if col_c.upper() not in GAS_LABELS:
            subcategory = col_c
            continue

        if col_c.upper() == "CH4" and category and gwp100 is not None:
            gwp20 = gwp20 if gwp20 is not None else 0.0
            records.append({
                "year": year,
                "category": category,
                "subcategory": subcategory or category,
                "mmtco2e_gwp20": gwp20,
                "mmtco2e_gwp100": gwp100,
                "ch4_tonnes": (gwp100 / GWP100_IN_WORKBOOK) * MMT_TO_TONNES,
            })

    return pd.DataFrame(records)


def build_rag_chunks(df: pd.DataFrame) -> list[dict]:
    """Natural-language summaries so the Chat's BM25 retrieval can surface these
    numbers — a raw CSV isn't searchable text, so it would never be retrieved."""
    chunks = []
    years = sorted(df["year"].unique())

    for year in years:
        sub = df[df["year"] == year]
        total_t = sub["ch4_tonnes"].sum()
        by_cat = (sub.groupby("category")["ch4_tonnes"].sum()
                    .sort_values(ascending=False))
        lines = ", ".join(f"{c} {v:,.0f} t" for c, v in by_cat.items() if v > 0)
        chunks.append({
            "text": (
                f"Maryland official state greenhouse gas inventory, {year}, methane (CH4). "
                f"Compiled by the Maryland Department of the Environment (MDE), not "
                f"Climate TRACE — this is Maryland's own bottom-up inventory. "
                f"Total CH4 across all categories: {total_t:,.0f} tonnes "
                f"({sub['mmtco2e_gwp100'].sum():.3f} MMTCO2e at AR5 100-year GWP, "
                f"{sub['mmtco2e_gwp20'].sum():.3f} MMTCO2e at 20-year GWP). "
                f"By category: {lines}."
            ),
            "source_file": f"MD_{year}_GHG_Inventory (MDE official)",
            "source_path": f"__md_inventory__/{year}",
            "iso3": "USA", "location": "Maryland",
            "tier": "smac_member", "sector": None,
            "output_types": ["data", "trend", "policy", "method"],
        })

        # per-category detail, so a question about landfills retrieves landfills
        for cat, grp in sub.groupby("category"):
            detail = ", ".join(
                f"{r.subcategory}: {r.ch4_tonnes:,.0f} t CH4"
                for r in grp.itertuples() if r.ch4_tonnes > 0
            )
            if not detail:
                continue
            chunks.append({
                "text": (
                    f"Maryland MDE official inventory {year} — {cat} methane detail. "
                    f"{detail}. Source: Maryland Department of the Environment state "
                    f"greenhouse gas inventory, converted from MMTCO2e at AR5 100-year "
                    f"GWP of {GWP100_IN_WORKBOOK:g}."
                ),
                "source_file": f"MD_{year}_GHG_Inventory (MDE official)",
                "source_path": f"__md_inventory__/{year}/{cat}",
                "iso3": "USA", "location": "Maryland",
                "tier": "smac_member", "sector": None,
                "output_types": ["data", "trend", "policy", "method"],
            })

    # a trend chunk across all years — the comparison that matters most
    totals = df.groupby("year")["ch4_tonnes"].sum().sort_index()
    if len(totals) > 1:
        trend = ", ".join(f"{y}: {v:,.0f} t" for y, v in totals.items())
        first, last = totals.iloc[0], totals.iloc[-1]
        pct = (last - first) / first * 100 if first else 0
        chunks.append({
            "text": (
                f"Maryland official MDE inventory, methane trend across inventory years. "
                f"{trend}. That is a {pct:+.1f}% change from {totals.index[0]} to "
                f"{totals.index[-1]} in Maryland's own bottom-up accounting. Note this is "
                f"a different data source from Climate TRACE, which this tool uses "
                f"elsewhere and which is modeled from satellite and activity data — the "
                f"two do not necessarily agree on level or on direction, and that "
                f"divergence is itself a documented finding for Maryland."
            ),
            "source_file": "MD GHG Inventory series (MDE official)",
            "source_path": "__md_inventory__/trend",
            "iso3": "USA", "location": "Maryland",
            "tier": "smac_member", "sector": None,
            "output_types": ["data", "trend", "policy", "method"],
        })
    return chunks


def main():
    files = sorted(UPLOAD_DIR.glob("MD_*.xlsx"))
    if not files:
        sys.exit(f"No MD_*.xlsx found in {UPLOAD_DIR.resolve()}\n"
                 f"Create that folder next to the repo root and put the five "
                 f"inventory workbooks in it.")

    frames = []
    for f in files:
        m = re.search(r"MD_(\d{4})_", f.name)
        if not m:
            print(f"  skip (no year in filename): {f.name}")
            continue
        year = int(m.group(1))
        df = parse_workbook(f, year)
        print(f"  {f.name}: {len(df)} CH4 rows, "
              f"{df['ch4_tonnes'].sum():,.0f} t CH4 total")
        frames.append(df)

    if not frames:
        sys.exit("Nothing parsed.")

    all_df = pd.concat(frames, ignore_index=True).sort_values(["year", "category"])
    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    all_df.to_csv(OUT_CSV, index=False)
    print(f"\nWrote {len(all_df)} rows -> {OUT_CSV}")
    print("\nCH4 tonnes by year:")
    print(all_df.groupby("year")["ch4_tonnes"].sum().round(0).to_string())

    # --- append to the RAG corpus (idempotent: strips any previous run first) ---
    chunks = build_rag_chunks(all_df)
    corpus = json.loads(RAG_PATH.read_text(encoding="utf-8")) if RAG_PATH.exists() else []
    before = len(corpus)
    corpus = [c for c in corpus
              if not str(c.get("source_path", "")).startswith("__md_inventory__")]
    stripped = before - len(corpus)
    next_id = max((c.get("id", 0) for c in corpus), default=-1) + 1
    for i, c in enumerate(chunks):
        c["id"] = next_id + i
    corpus.extend(chunks)
    RAG_PATH.write_text(json.dumps(corpus), encoding="utf-8")
    print(f"\nRAG corpus: removed {stripped} stale MD-inventory chunks, "
          f"added {len(chunks)} -> {len(corpus)} total")


if __name__ == "__main__":
    main()
